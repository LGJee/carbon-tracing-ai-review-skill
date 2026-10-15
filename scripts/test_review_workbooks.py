"""Local regression tests using synthetic workbooks, not biochemical evidence."""

import argparse
from contextlib import redirect_stdout
from io import StringIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.dont_write_bytecode = True
import review_workbooks as review

MAPPED = "[CH3:1][CH2:2][OH:3]>>[CH3:1][CH2:2][OH:3]"
SWAPPED = "[CH3:1][CH2:2][OH:3]>>[CH3:2][CH2:1][OH:3]"
PROJECT = None
MAPPING_SKILL = None


def fixture(path):
    workbook = Workbook()
    del workbook[workbook.active.title]
    equations = [
        "A[e] -> A[c]",
        "A[c] -> B[c]",
        "B[c] -> C[c]",
        "C[c] -> D[c]",
        "D[c] -> E[c]",
    ]
    ids = [0, "r1", "r2", "r3", "r4"]
    sheet = workbook.create_sheet("Nonzero")
    sheet.append(["ReactionID", "ReactionName", "Equation", "Flux"])
    for key, equation in zip(ids, equations):
        sheet.append([key, "Synthetic fixture", equation, 1])
    sheet = workbook.create_sheet("Metabolites")
    sheet.append(["MetaboliteID", "MetaboliteName", "SMILES_Model", "SMILES_Manual"])
    for i, name in enumerate(["A[e]", "A[c]", "B[c]", "C[c]", "D[c]", "E[c]"]):
        sheet.append(["m" + str(i), name, "CCO", None])
    for backend in review.BACKENDS:
        sheet = workbook.create_sheet(backend)
        sheet.append(
            [
                "ReactionID",
                "ReactionName",
                "Equation",
                "ReactionSMILES",
                "MappedReactionSMILES",
                "MappingDirection",
                "MappingConfident",
            ]
        )
        for key, equation in zip(ids, equations):
            sheet.append(
                [
                    key,
                    "Synthetic fixture",
                    equation,
                    "CCO>>CCO",
                    MAPPED,
                    "forward",
                    0.8273648273648273,
                ]
            )
    selection = workbook.create_sheet("Mapping_Selection")
    selection.append(
        [
            "ReactionID",
            "ReactionName",
            "Flux",
            "SelectedMapper",
            "ReviewRequired",
            "SelectionReason",
            "Equation",
            "MappingDirection",
            "SelectedMappedReactionSMILES",
        ]
    )
    final = workbook.create_sheet("Final_Mapping")
    final.append(
        [
            "ReactionID",
            "ReactionName",
            "Flux",
            "SelectedMapper",
            "ReviewRequired",
            "AIGenerated",
            "MappedReactionSMILES",
            "Equation",
            "MappingDirection",
        ]
    )
    for key, equation, flag in zip(
        ids, equations, [False, "TRUE", True, "FALSE", False]
    ):
        selection.append(
            [
                key,
                "Synthetic fixture",
                1,
                "LocalMapper",
                flag,
                "Fixture only",
                equation,
                "forward",
                None if key == "r2" else MAPPED,
            ]
        )
        if key != "r2":
            final.append(
                [
                    key,
                    "Synthetic fixture",
                    1,
                    "LocalMapper",
                    key in ("r1", "r3"),
                    False,
                    MAPPED,
                    equation,
                    "forward",
                ]
            )
    info = workbook.create_sheet("Mapping_Info")
    info.append(["Mapper", "SourceInputSHA256", "OptionsJSON"])
    info.append(["LocalMapper", "a" * 64, "{}"])
    info.sheet_state = "hidden"
    workbook.save(path)
    workbook.close()


def proposal(key, mapped=MAPPED):
    return dict(
        ReactionID=key,
        MappedReactionSMILES=mapped,
        scope="all",
        evidence_status="evidence-supported",
        rationale="Synthetic test fixture; not a scientific evidence claim.",
        limitations=["Fixture only"],
        evidence=[
            dict(
                url="https://example.org/fixture",
                accessed_on="2026-10-14",
                claim="Synthetic test metadata, not retrieved evidence.",
            )
        ],
    )


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="carbon-review-test-")
        self.folder = Path(self.temporary.name)
        self.source = self.folder / "Mixed_Culture_2_atom_mapping.xlsx"
        fixture(self.source)
        self.before = review.digest(self.source)
        # Fixtures must never enter the real project's scientific review archive.
        writer = review.archive_review

        def save_archive(report, targets, proposals, directory):
            if directory == PROJECT / "data/ai_mapping_reviews":
                directory = self.folder / "project/data/ai_mapping_reviews"
            return writer(report, targets, proposals, directory)

        patcher = patch.object(review, "archive_review", side_effect=save_archive)
        patcher.start()
        self.addCleanup(patcher.stop)

    def tearDown(self):
        self.assertEqual(review.digest(self.source), self.before)
        self.temporary.cleanup()

    def entry(self, replacements):
        return dict(
            input_file=str(self.source),
            input_sha256=self.before,
            replacements=replacements,
        )

    def apply(self, replacements, use_project=True, output="reviewed"):
        return review.apply_workbook(
            self.entry(replacements),
            PROJECT if use_project else None,
            MAPPING_SKILL,
            self.folder / output,
        )

    def test_flags_and_union_include_missing_final_rows(self):
        self.assertFalse(review.flag(" FALSE "))
        self.assertTrue(review.flag("true"))
        self.assertEqual(review.identity(0), "0")
        self.assertEqual(review.identity(12.0), "12")
        with self.assertRaises(ValueError):
            review.flag("Ture")
        with self.assertRaises(ValueError):
            review.flag(1)
        workbook = load_workbook(self.source)
        try:
            targets, issues = review.review_targets(workbook)
        finally:
            workbook.close()
        self.assertEqual(set(targets), {"r1", "r2", "r3"})
        self.assertTrue(targets["r2"]["missing_final_row"])
        self.assertFalse(issues)

    def test_success_insertion_failure_and_provenance_preservation(self):
        report = self.apply(
            [
                proposal("r1", SWAPPED),
                proposal("r2"),
                proposal("r3", "[CH4:1]>>[CH4:1]"),
            ]
        )
        self.assertEqual(report["accepted"], 2)
        output = Path(report["output_file"])
        original, reviewed = load_workbook(self.source), load_workbook(output)
        try:
            for name in [
                "Nonzero",
                "Metabolites",
                "Mapping_Info",
                "Mapping_Selection",
            ] + list(review.BACKENDS):
                self.assertEqual(
                    list(original[name].values), list(reviewed[name].values)
                )
            self.assertEqual(reviewed["Mapping_Info"].sheet_state, "hidden")
            _, final = review.table(reviewed, "Final_Mapping", review.FINAL_COLUMNS)
            _, selection = review.table(
                reviewed, "Mapping_Selection", review.SELECTION_COLUMNS
            )
            self.assertEqual(len(final), 5)
            for key in ("r1", "r2"):
                self.assertIs(final[key]["AIGenerated"], True)
                self.assertIs(final[key]["AIReviewed"], True)
                self.assertEqual(final[key]["SelectedMapper"], "AI-assisted")
            self.assertEqual(final["r1"]["MappedReactionSMILES"], SWAPPED)
            self.assertEqual(selection["r1"]["SelectedMappedReactionSMILES"], MAPPED)
            self.assertIsNone(selection["r2"]["SelectedMappedReactionSMILES"])
            self.assertEqual(final["r2"]["MappedReactionSMILES"], MAPPED)
            self.assertIs(final["r1"]["ReviewRequired"], True)
            self.assertIs(final["r2"]["ReviewRequired"], True)
            self.assertIs(final["r3"]["AIGenerated"], False)
            self.assertIs(final["r3"]["AIReviewed"], False)
            self.assertEqual(final["r3"]["MappedReactionSMILES"], MAPPED)
            self.assertEqual(selection["r4"]["SelectedMappedReactionSMILES"], MAPPED)
        finally:
            original.close()
            reviewed.close()

    def test_hash_and_existing_completed_output_are_protected(self):
        wrong = self.entry([proposal("r1")])
        wrong["input_sha256"] = "0" * 64
        with self.assertRaisesRegex(ValueError, "SHA256"):
            review.apply_workbook(
                wrong, PROJECT, MAPPING_SKILL, self.folder / "reviewed"
            )
        report = self.apply([proposal("r1")])
        output = Path(report["output_file"])
        saved = review.digest(output)
        with self.assertRaisesRegex(ValueError, "retained"):
            self.apply([proposal("r1", SWAPPED)])
        self.assertEqual(review.digest(output), saved)

    def test_invalid_scope_evidence_map_and_unflagged_rows_are_unchanged(self):
        entries = [proposal("r1"), proposal("r2"), proposal("r3"), proposal("r4")]
        entries[0]["scope"] = "carbon"
        entries[1]["evidence_status"] = "unresolved"
        entries[2]["MappedReactionSMILES"] = (
            "[CH3:1][CH2:1][OH:3]>>[CH3:1][CH2:2][OH:3]"
        )
        report = self.apply(entries)
        self.assertEqual(report["accepted"], 0)
        self.assertIsNone(report["output_file"])
        self.assertTrue(all(row["status"] == "rejected" for row in report["reactions"]))

    def test_raw_input_mismatch_and_duplicate_proposals_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "Duplicated proposal"):
            self.apply([proposal("r1"), proposal("r1")])
        workbook = load_workbook(self.source)
        workbook["RXNMapper"]["D3"] = "C>>C"
        workbook.save(self.source)
        workbook.close()
        self.before = review.digest(self.source)
        report = self.apply([proposal("r1")])
        self.assertEqual(report["accepted"], 0)

    def test_no_project_still_runs_scientific_structural_validation(self):
        report = self.apply([proposal("r1")], use_project=False)
        self.assertEqual(report["accepted"], 1)
        self.assertEqual(report["project_validation"], "not_available")

    def test_invalid_evidence_and_cell_overflow_cannot_change_failed_rows(self):
        bad_evidence = proposal("r2")
        bad_evidence["evidence"][0]["url"] = 42
        bad_map = proposal("r1", "C" * 32768)
        report = self.apply([bad_evidence, bad_map])
        self.assertEqual(report["accepted"], 0)
        too_long = proposal("r3")
        too_long["rationale"] = "x" * 32768
        report = self.apply([proposal("r1"), too_long])
        self.assertEqual(report["accepted"], 2)
        reviewed = load_workbook(report["output_file"])
        original = load_workbook(self.source)
        try:
            _, before = review.table(
                original, "Mapping_Selection", review.SELECTION_COLUMNS
            )
            _, after = review.table(
                reviewed, "Mapping_Selection", review.SELECTION_COLUMNS
            )
            self.assertEqual(before["r3"], after["r3"])
        finally:
            reviewed.close()
            original.close()

    def test_numeric_zero_id_is_not_dropped(self):
        workbook = load_workbook(self.source)
        sheet = workbook["Mapping_Selection"]
        column = [c.value for c in sheet[1]].index("ReviewRequired") + 1
        sheet.cell(2, column, True)
        workbook.save(self.source)
        workbook.close()
        self.before = review.digest(self.source)
        report = self.apply([proposal(0)])
        self.assertEqual(report["accepted"], 1)
        reviewed = load_workbook(report["output_file"])
        try:
            _, rows = review.table(reviewed, "Final_Mapping", review.FINAL_COLUMNS)
            self.assertIs(rows["0"]["AIGenerated"], True)
        finally:
            reviewed.close()

    def test_multiple_workbooks_continue_after_one_failure(self):
        second = self.folder / "Other_Culture_2_atom_mapping.xlsx"
        fixture(second)
        invalid = self.entry([proposal("r1")])
        invalid["input_sha256"] = "0" * 64
        valid = dict(
            input_file=str(second),
            input_sha256=review.digest(second),
            replacements=[proposal("r1")],
        )
        proposals = self.folder / "proposals.json"
        proposals.write_text(
            json.dumps(dict(workbooks=[invalid, valid])), encoding="utf-8"
        )
        stream = StringIO()
        arguments = [
            "apply",
            str(proposals),
            "--project",
            str(PROJECT),
            "--mapping-skill",
            str(MAPPING_SKILL),
            "--output-dir",
            str(self.folder / "out"),
        ]
        with redirect_stdout(stream):
            result = review.main(arguments)
        report = json.loads(stream.getvalue())["workbooks"]
        self.assertEqual(result, 2)
        self.assertIn("error", report[0])
        self.assertEqual(report[1]["accepted"], 1)
        self.assertTrue(Path(report[1]["output_file"]).is_file())
        self.assertEqual(review.digest(second), valid["input_sha256"])
        delivery = json.loads(stream.getvalue())
        self.assertTrue(Path(delivery["report_files"]["markdown"]).is_file())
        saved = json.loads(Path(delivery["report_files"]["json"]).read_text())
        self.assertEqual(saved["workbooks"], report)

    def test_project_rejects_mismatched_equation_before_writing(self):
        workbook = load_workbook(self.source)
        for name in list(review.BACKENDS) + ["Mapping_Selection", "Final_Mapping"]:
            sheet = workbook[name]
            column = [c.value for c in sheet[1]].index("MappingDirection") + 1
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == "r1":
                    sheet.cell(row, column, "reverse")
        workbook.save(self.source)
        workbook.close()
        self.before = review.digest(self.source)
        report = self.apply([proposal("r1")])
        self.assertEqual(report["accepted"], 0)

    def test_written_copy_is_consumed_by_existing_step4_path(self):
        report = self.apply([proposal("r1", SWAPPED), proposal("r2")])
        from carbon_flow_tools import pipeline
        from carbon_flow_tools.cli import main
        from carbon_flow_tools.atom_helpers import assignment_from_labeled_smiles

        options = dict(
            ai_generated=True,
            input_dir=str(self.folder / "reviewed"),
            output_dir=str(self.folder / "reviewed"),
            rounds=1,
            carbon_sources={"Mixed_Culture": ["A[e]"]},
            max_carbon_processes=1,
            carbon_memory_gb=None,
            tracing_python=sys.executable,
        )
        with patch.object(
            pipeline, "run_mapper", side_effect=AssertionError("Mapper invoked")
        ):
            with redirect_stdout(StringIO()):
                self.assertEqual(main("trace", options, []), 0)
        output = self.folder / "reviewed/Mixed_Culture_3_carbon_flow_AIGenerated.xlsx"
        workbook = load_workbook(output)
        try:
            self.assertEqual(workbook.sheetnames, ["Summary", "Carbon_Flow_R1"])
            records = list(workbook["Carbon_Flow_R1"].values)
            rows = [dict(zip(records[0], row)) for row in records[1:]]
            product = next(row for row in rows if row["Production"] == "B[c]")
            self.assertEqual(
                assignment_from_labeled_smiles("CCO", product["LabeledSMILES"]),
                ((0, 0, "C2"), (0, 1, "C1")),
            )
            self.assertIn("C[c]", {row["Production"] for row in rows})
        finally:
            workbook.close()
        self.assertTrue(Path(report["output_file"]).is_file())

    def test_archive_retains_inputs_evidence_and_failed_proposals(self):
        report = self.apply(
            [proposal("r1"), proposal("r2", "[CH4:1]>>[CH4:1]")], use_project=False
        )
        path = Path(report["archive_file"])
        self.assertEqual(path.parent, self.folder / "reviewed/data/ai_mapping_reviews")
        archive = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(archive["input_sha256"], self.before)
        records = {row["target"]["ReactionID"]: row for row in archive["reactions"]}
        self.assertEqual(records["r1"]["target"]["ReactionSMILES"], "CCO>>CCO")
        self.assertEqual(
            records["r1"]["proposal"]["evidence"], proposal("r1")["evidence"]
        )
        self.assertEqual(
            records["r1"]["outcome"]["checks"][0]["status"], "structural_checks_passed"
        )
        self.assertEqual(records["r2"]["outcome"]["status"], "rejected")
        book = load_workbook(self.source)
        try:
            targets, _ = review.review_targets(book)
        finally:
            book.close()
        candidates, issues = review.cached_reviews(self.folder / "reviewed", targets)
        self.assertFalse(issues)
        self.assertEqual(set(candidates), {"r1"})
        self.assertTrue(candidates["r1"][0]["requires_revalidation"])
        self.assertEqual(
            candidates["r1"][0]["proposal"]["MappedReactionSMILES"], MAPPED
        )
        cached = candidates["r1"][0]["proposal"]
        cached["MappedReactionSMILES"] = "[CH4:1]>>[CH4:1]"
        with self.assertRaises(ValueError):
            review.check_proposal(
                cached,
                targets["r1"],
                review.scientific_validator(MAPPING_SKILL),
                {},
                None,
            )
        targets["r1"]["MappingDirection"] = "reverse"
        self.assertEqual(
            review.cached_reviews(self.folder / "reviewed", targets)[0], {}
        )
        targets["r1"]["MappingDirection"] = "forward"
        targets["r1"]["input_variants"] = ["CC[O-]>>CC[O-]"]
        self.assertEqual(
            review.cached_reviews(self.folder / "reviewed", targets)[0], {}
        )

    def test_rejected_archive_is_not_reused_and_bad_json_is_reported(self):
        report = self.apply([proposal("r1", "[CH4:1]>>[CH4:1]")], use_project=False)
        self.assertIsNone(report["output_file"])
        self.assertTrue(Path(report["archive_file"]).is_file())
        book = load_workbook(self.source)
        try:
            targets, _ = review.review_targets(book)
        finally:
            book.close()
        directory = self.folder / "reviewed/data/ai_mapping_reviews"
        (directory / "broken.json").write_text("{", encoding="utf-8")
        candidates, issues = review.cached_reviews(self.folder / "reviewed", targets)
        self.assertEqual(candidates, {})
        self.assertEqual(len(issues), 1)

    def test_archive_write_failure_reports_but_retains_delivered_workbook(self):
        with patch.object(
            review, "archive_review", side_effect=PermissionError("read-only data")
        ):
            report = self.apply([proposal("r1")], use_project=False)
        self.assertTrue(Path(report["output_file"]).is_file())
        self.assertIn("read-only data", report["archive_error"])

    def test_existing_review_archives_are_preserved(self):
        first = self.apply([proposal("r1")], output="reviewed")
        first_path = Path(first["archive_file"])
        before = review.digest(first_path)
        second = self.apply([proposal("r1", SWAPPED)], output="second")
        self.assertNotEqual(first["archive_file"], second["archive_file"])
        self.assertEqual(first_path.parent, Path(second["archive_file"]).parent)
        self.assertEqual(review.digest(first_path), before)

    def test_review_only_distinguishes_retained_unresolved_and_screening(self):
        entry = self.entry([])
        entry["reviews"] = [
            dict(
                ReactionID="r1",
                decision="unchanged",
                AIReviewed=True,
                rationale="Scientific review retained the existing map.",
            ),
            dict(
                ReactionID="r3",
                decision="unresolved",
                AIReviewed=True,
                rationale="Scientific analysis did not resolve the mechanism.",
            ),
            dict(
                ReactionID="r2",
                decision="unresolved",
                AIReviewed=False,
                rationale="Input screening only; scientific review not performed.",
            ),
        ]
        report = review.apply_workbook(
            entry, PROJECT, MAPPING_SKILL, self.folder / "out"
        )
        self.assertEqual(report["accepted"], 0)
        self.assertEqual(report["ai_reviewed"], 2)
        self.assertEqual(report["project_validation"], "not_run")
        workbook = load_workbook(report["output_file"])
        try:
            columns, rows = review.table(
                workbook, "Final_Mapping", review.FINAL_COLUMNS
            )
            self.assertEqual(columns[columns.index("ReviewRequired") + 1], "AIReviewed")
            self.assertNotIn("r2", rows)
            for key in ("r1", "r3"):
                self.assertIs(rows[key]["AIReviewed"], True)
                self.assertIs(rows[key]["AIGenerated"], False)
                self.assertEqual(rows[key]["MappedReactionSMILES"], MAPPED)
        finally:
            workbook.close()
        paths = review.write_delivery_report([report], self.folder / "out")
        text = Path(paths["markdown"]).read_text(encoding="utf-8")
        self.assertIn("| Total | 3 | 2 | 0 | 1 | 2 | 0 |", text)
        self.assertEqual(
            len(json.loads(Path(paths["json"]).read_text())["workbooks"][0]["reviews"]),
            3,
        )
        archive = json.loads(Path(report["archive_file"]).read_text())
        self.assertEqual(archive["reviews"], entry["reviews"])

    def test_bad_review_metadata_cannot_mark_a_row(self):
        for changes in (
            {"AIReviewed": "TRUE"},
            {"decision": "correct"},
            {"ReactionID": "r4"},
            {"rationale": ""},
        ):
            entry = self.entry([])
            row = dict(
                ReactionID="r1",
                decision="unresolved",
                AIReviewed=True,
                rationale="Incomplete evidence after analysis.",
            )
            row.update(changes)
            entry["reviews"] = [row]
            with self.assertRaises(ValueError):
                review.apply_workbook(
                    entry, PROJECT, MAPPING_SKILL, self.folder / "bad"
                )
        self.assertFalse((self.folder / "bad").exists())

    def test_existing_review_column_is_moved_without_losing_flags(self):
        workbook = load_workbook(self.source)
        try:
            sheet = workbook["Final_Mapping"]
            column = sheet.max_column + 1
            sheet.cell(1, column, "AIReviewed")
            for row in range(2, sheet.max_row + 1):
                sheet.cell(row, column, row == 2)
            review.mark_reviewed(workbook, {"r3"})
            headers, records = review.table(
                workbook, "Final_Mapping", review.FINAL_COLUMNS
            )
            self.assertEqual(headers.count("AIReviewed"), 1)
            self.assertEqual(
                headers.index("AIReviewed"), headers.index("ReviewRequired") + 1
            )
            self.assertIs(records["0"]["AIReviewed"], True)
            self.assertIs(records["r3"]["AIReviewed"], True)
            self.assertIs(records["r1"]["AIReviewed"], False)
        finally:
            workbook.close()

    def test_report_separates_unassessed_and_validation_rejections(self):
        report = self.apply([proposal("r1", "[CH4:1]>>[CH4:1]"), proposal("r4")])
        paths = review.write_delivery_report([report], self.folder / "reports")
        text = Path(paths["markdown"]).read_text(encoding="utf-8")
        self.assertIn("| Total | 3 | 0 | 0 | 0 | 1 | 2 |", text)
        second = review.write_delivery_report([report], self.folder / "reports")
        self.assertNotEqual(paths, second)
        self.assertTrue(Path(paths["markdown"]).is_file())


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--mapping-skill", type=Path, required=True)
    args, other = parser.parse_known_args()
    PROJECT = args.project.resolve()
    MAPPING_SKILL = args.mapping_skill.resolve()
    unittest.main(argv=[sys.argv[0]] + other)
